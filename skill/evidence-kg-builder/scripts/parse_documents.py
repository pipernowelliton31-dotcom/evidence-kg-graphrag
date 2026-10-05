"""Lightweight provenance-preserving parsers for heterogeneous business documents."""
from __future__ import annotations

import copy
import hashlib
import importlib.metadata
import io
import re
import shutil
import zipfile
from datetime import date, datetime
from pathlib import Path

from kg_core import (CONFIG, ENGINE_VERSION, SCHEMA_VERSION, ContractError, canonical,
                     clean, empty_alignment, engine_signature, ident, read_json, require, save_state, write_json)

OFFICE = {".docx", ".pptx", ".xlsx"}
SUPPORTED = OFFICE | {".pdf", ".png", ".jpg", ".jpeg", ".webp", ".md", ".txt"}


def hash_file(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def dependency_versions():
    packages = ["openpyxl", "python-pptx", "pdfplumber", "Pillow", "defusedxml",
                "pdfminer-six", "pypdfium2", "requests", "python-dotenv"]
    result = {}
    for package in packages:
        try:
            result[package] = importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError:
            result[package] = "missing"
    return result


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return clean(value)


def md_table(rows):
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    def row_line(row):
        values = [cell_text(x).replace("|", "\\|").replace("\n", "<br>") for x in row]
        values += [""] * (width - len(values))
        return "| " + " | ".join(values) + " |"
    return "\n".join([row_line(rows[0]), row_line(["---"] * width), *[row_line(r) for r in rows[1:]]])


def _pdf_same_x_span(a, b, tolerance=18.0):
    if not a or not b or len(a) != 4 or len(b) != 4:
        return False
    try:
        return abs(float(a[0]) - float(b[0])) <= tolerance and abs(float(a[2]) - float(b[2])) <= tolerance
    except (TypeError, ValueError):
        return False


def _pdf_words_in_bbox(words, bbox, pad=1.5):
    x0, top, x1, bottom = map(float, bbox)
    return [w for w in words
            if float(w["x1"]) >= x0 - pad and float(w["x0"]) <= x1 + pad
            and float(w["bottom"]) >= top - pad and float(w["top"]) <= bottom + pad]


def _pdf_merge_cell_words(words):
    if not words:
        return ""
    ordered = sorted(words, key=lambda w: (float(w["top"]), float(w["x0"])))
    lines = []
    for word in ordered:
        top = float(word["top"])
        if not lines or abs(top - lines[-1][0]) > 3.0:
            lines.append([top, [word]])
        else:
            lines[-1][1].append(word)
    rendered = []
    for _, line_words in lines:
        line_words.sort(key=lambda w: float(w["x0"]))
        text, previous = "", None
        for word in line_words:
            token = str(word["text"])
            if previous is not None:
                gap = float(word["x0"]) - float(previous["x1"])
                if (gap > 2.5 and re.search(r"[A-Za-z0-9]$", str(previous["text"]))
                        and re.match(r"^[A-Za-z0-9]", token)):
                    text += " "
            text += token
            previous = word
        rendered.append(text)
    return "\n".join(rendered)


def _pdf_infer_header_schema(words, row_bbox, table_bbox, page_number):
    header_words = _pdf_words_in_bbox(words, row_bbox)
    if len(header_words) < 2:
        return None
    clusters = []
    for word in sorted(header_words, key=lambda w: (float(w["x0"]), float(w["top"]))):
        wx0, wx1 = float(word["x0"]), float(word["x1"])
        best, best_score = None, -1e9
        for cluster in clusters:
            overlap = min(cluster["x1"], wx1) - max(cluster["x0"], wx0)
            gap = max(cluster["x0"] - wx1, wx0 - cluster["x1"], 0)
            score = overlap if overlap >= 0 else -gap
            if (overlap >= -2.0 or gap <= 8.0) and score > best_score:
                best, best_score = cluster, score
        if best is None:
            clusters.append({"x0": wx0, "x1": wx1, "words": [word]})
        else:
            best["x0"] = min(best["x0"], wx0)
            best["x1"] = max(best["x1"], wx1)
            best["words"].append(word)
    clusters.sort(key=lambda c: c["x0"])
    clusters = [c for c in clusters
                if re.search(r"[A-Za-z0-9\u4e00-\u9fff]", _pdf_merge_cell_words(c["words"]))]
    if len(clusters) < 2:
        return None
    x0, _, x1, _ = map(float, table_bbox)
    cuts = [x0]
    for left, right in zip(clusters, clusters[1:]):
        cuts.append((left["x1"] + right["x0"]) / 2.0)
    cuts.append(x1)
    header = [_pdf_merge_cell_words(c["words"]).replace("\n", "") for c in clusters]
    bounds = [(cuts[i], cuts[i + 1]) for i in range(len(header))]
    return {"page": page_number, "header": header, "bounds": bounds,
            "bbox": list(map(float, table_bbox))}


def _pdf_row_from_schema(words, row_bbox, schema):
    row_words = _pdf_words_in_bbox(words, row_bbox)
    cells = []
    last_right = schema["bounds"][-1][1]
    for left, right in schema["bounds"]:
        selected = []
        for word in row_words:
            center = (float(word["x0"]) + float(word["x1"])) / 2.0
            in_band = left <= center < right or (right == last_right and left <= center <= right)
            if in_band:
                selected.append(word)
        cells.append(_pdf_merge_cell_words(selected))
    return cells


def _pdf_reconstruct_table(words, table, previous_schema, page_number):
    rows = list(getattr(table, "rows", []) or [])
    if not rows:
        raw = table.extract() or []
        return (raw[0], raw[1:], None) if raw else ([], [], None)
    if (previous_schema and previous_schema.get("bounds") and previous_schema.get("page") == page_number - 1
            and _pdf_same_x_span(previous_schema.get("bbox"), table.bbox)):
        first = _pdf_row_from_schema(words, rows[0].bbox, previous_schema)
        old = [norm for norm in (re.sub(r"\s+", "", x) for x in previous_schema["header"])]
        now = [norm for norm in (re.sub(r"\s+", "", x) for x in first)]
        matches = sum(bool(a and b and (a == b or a in b or b in a)) for a, b in zip(old, now))
        repeated_header = matches >= max(2, len(old) // 2)
        start = 1 if repeated_header else 0
        data_rows = [_pdf_row_from_schema(words, row.bbox, previous_schema) for row in rows[start:]]
        schema = {**previous_schema, "page": page_number, "bbox": list(map(float, table.bbox)),
                  "inherited_from_page": previous_schema["page"]}
        return list(previous_schema["header"]), data_rows, schema
    schema = _pdf_infer_header_schema(words, rows[0].bbox, table.bbox, page_number)
    if schema:
        data_rows = [_pdf_row_from_schema(words, row.bbox, schema) for row in rows[1:]]
        return list(schema["header"]), data_rows, schema
    raw = table.extract() or []
    if not raw:
        return [], [], None
    return raw[0], raw[1:], {"page": page_number, "header": raw[0],
                             "bbox": list(map(float, table.bbox))}


def split_text(text, budget):
    """Split on paragraphs/lines; only oversized plain prose is split at sentence boundaries."""
    result, current = [], ""
    for block in re.split(r"(\n\s*\n)", clean(text)):
        if not block:
            continue
        if len(current) + len(block) > budget and current.strip():
            result.append(current.strip())
            current = ""
        if len(block) > budget:
            lines = block.splitlines(keepends=True)
            for line in lines:
                if len(current) + len(line) > budget and current.strip():
                    result.append(current.strip())
                    current = ""
                if len(line) > budget and not line.lstrip().startswith("|"):
                    parts = re.split(r"(?<=[。！？.!?])", line)
                    for part in parts:
                        if len(current) + len(part) > budget and current.strip():
                            result.append(current.strip())
                            current = ""
                        current += part
                else:
                    current += line
        else:
            current += block
    if current.strip():
        result.append(current.strip())
    return result or [""]


class Adapter:
    def __init__(self, document, path, work, config):
        self.document, self.path, self.work, self.config = document, path, Path(work), config
        self.units = []
        self.assets_dir = self.work / "assets" / document["document_id"]
        self.assets_dir.mkdir(parents=True, exist_ok=True)

    def asset(self, name, data):
        path = self.assets_dir / Path(name).name
        path.write_bytes(data)
        return path.relative_to(self.work).as_posix()

    def add(self, content, locator, assets=None, warnings=None, structured=None, split=True):
        blocks = split_text(content, self.config["max_chars"]) if split else [clean(content)]
        for i, block in enumerate(blocks):
            location = copy.deepcopy(locator)
            if len(blocks) > 1:
                location["part"] = i + 1
            warning = list(warnings or [])
            if len(block) > self.config["max_chars"]:
                warning.append("oversized_atomic_block: inspect without truncating cells")
            unit = {"document_id": self.document["document_id"], "format": self.document["format"],
                    "locator": location, "content_md": block, "asset_refs": list(assets or []),
                    "parse_warnings": warning}
            if structured is not None:
                unit["structured"] = structured
            unit["unit_id"] = ident("unit", [self.document["document_id"], location, block, structured])
            self.units.append(unit)


    def docx(self):
        from defusedxml import ElementTree as ET
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
              "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        with zipfile.ZipFile(self.path) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
            style_headings = {}
            if "word/styles.xml" in archive.namelist():
                for style in ET.fromstring(archive.read("word/styles.xml")).findall("w:style", ns):
                    style_id = style.get("{" + ns["w"] + "}styleId")
                    name = style.find("w:name", ns)
                    outline = style.find("w:pPr/w:outlineLvl", ns)
                    value = name.get("{" + ns["w"] + "}val", "") if name is not None else ""
                    if outline is not None:
                        style_headings[style_id] = int(outline.get("{" + ns["w"] + "}val", "0")) + 1
                    elif re.search(r"heading|标题", value, re.I):
                        number = re.search(r"\d+", value)
                        style_headings[style_id] = int(number.group()) if number else 1
            section = []
            buffer, start, last = [], None, None
            para_index, table_index = 0, 0
            def flush():
                nonlocal buffer, start, last
                if buffer:
                    self.add("\n\n".join(buffer), {"section": list(section), "paragraph_start": start,
                             "paragraph_end": last, "precision": "source_xml_blocks"})
                buffer, start, last = [], None, None
            for block in root.find("w:body", ns):
                tag = block.tag.rsplit("}", 1)[-1]
                if tag == "p":
                    para_index += 1
                    text = "".join((x.text or "") for x in block.findall(".//w:t", ns))
                    style = block.find("w:pPr/w:pStyle", ns)
                    sid = style.get("{" + ns["w"] + "}val", "") if style is not None else ""
                    level = style_headings.get(sid)
                    if level and text.strip():
                        flush()
                        section = section[:level - 1] + [text]
                        text = "#" * min(level, 6) + " " + text
                    if text.strip():
                        if start is None:
                            start = para_index
                        last = para_index
                        buffer.append(text)
                        if sum(len(x) for x in buffer) > self.config["max_chars"]:
                            flush()
                    if block.findall(".//w:drawing", ns):
                        flush()
                        self.add("嵌入图形：需检查原始 DOCX 图形内容。", {"section": list(section), "paragraph": para_index},
                                 warnings=["docx_embedded_drawing_requires_host_review"])
                elif tag == "tbl":
                    flush()
                    table_index += 1
                    records = []
                    for ri, row in enumerate(block.findall("w:tr", ns), 1):
                        cells = []
                        for ci, cell in enumerate(row.findall("w:tc", ns), 1):
                            text = "\n".join("".join(x.text or "" for x in p.findall(".//w:t", ns)) for p in cell.findall("w:p", ns))
                            span = cell.find("w:tcPr/w:gridSpan", ns)
                            merge = cell.find("w:tcPr/w:vMerge", ns)
                            cells.append({"row": ri, "cell": ci, "text": text,
                                "grid_span": int(span.get("{" + ns["w"] + "}val", "1")) if span is not None else 1,
                                "vertical_merge": (merge.get("{" + ns["w"] + "}val", "continue") if merge is not None else None)})
                        records.append(cells)
                    if records:
                        header = [c["text"] for c in records[0]]
                        self.table_batches(records[1:], header, {"section": list(section), "table": table_index},
                            lambda row: [c["text"] for c in row], "source_xml_table", first_record=2,
                            extra={"header_cells": records[0]})
            flush()
            # All embedded raster images remain host tasks, even if not captioned.
            for name in archive.namelist():
                if name.startswith("word/media/") and Path(name).suffix.lower() in SUPPORTED - OFFICE - {".pdf", ".md", ".txt"}:
                    ref = self.asset(Path(name).name, archive.read(name))
                    self.add("DOCX embedded image: " + Path(name).name, {"media_part": name}, assets=[ref],
                             warnings=["host_visual_review_required"])
        if not self.units:
            self.add("[DOCX contains no extractable text/table blocks]", {"section": [], "precision": "source_xml_blocks"},
                     warnings=["docx_text_blocks_unavailable: inspect embedded assets if present"])

    def table_batches(self, records, header, locator, display, kind, first_record=1, extra=None):
        current, rendered_size, start = [], 0, first_record
        batches = []
        overhead = len(md_table([header])) + len(canonical(extra or {}))
        for offset, record in enumerate(records):
            size = sum(len(cell_text(x)) for x in display(record)) + min(256, len(canonical(record)) // 12)
            if current and (len(current) >= self.config["max_rows"] or rendered_size + size + overhead > self.config["max_chars"]):
                batches.append((start, current))
                current, rendered_size, start = [], 0, first_record + offset
            current.append(record)
            rendered_size += size
        if current:
            batches.append((start, current))
        if not batches:
            batches = [(first_record, [])]
        for start, batch in batches:
            location = {**locator, "row_start": start, "row_end": start + len(batch) - 1}
            content = md_table([header, *[display(r) for r in batch]])
            self.add(content, location, structured={"kind": kind, "records": batch, **(extra or {})}, split=False)

    def xlsx(self):
        import openpyxl
        # XLSX already has a precise openpyxl ledger; avoid a redundant MarkItDown pass.
        md_warning = []
        formulas = openpyxl.load_workbook(self.path, data_only=False)
        cached = openpyxl.load_workbook(self.path, data_only=True)
        try:
            for sheet in formulas:
                values = cached[sheet.title]
                merged = [str(r) for r in sheet.merged_cells.ranges]
                # Preserve header candidate rows verbatim; the host decides header semantics.
                header_rows = []
                for row in sheet.iter_rows(min_row=1, max_row=min(3, sheet.max_row)):
                    header_rows.append([{ "coordinate": c.coordinate, "raw": cell_text(c.value)} for c in row])
                records = []
                for ri, row in enumerate(sheet.iter_rows(), 1):
                    cells = []
                    for c in row:
                        if c.value is None and not c.has_style:
                            continue
                        formula = c.value if c.data_type == "f" else None
                        cells.append({"coordinate": c.coordinate, "raw": cell_text(c.value),
                                      "value": cell_text(values[c.coordinate].value) if formula else cell_text(c.value),
                                      "formula": formula, "cached_value": cell_text(values[c.coordinate].value) if formula else None,
                                      "comment": c.comment.text if c.comment else None,
                                      "number_format": c.number_format})
                    if cells:
                        records.append({"row": ri, "hidden": bool(sheet.row_dimensions[ri].hidden), "cells": cells})
                headings = ["Excel row", "Hidden", "Cell", "Original / cached value", "Formula / comment"]
                def display(record):
                    return [record["row"], str(record["hidden"]),
                            "; ".join(c["coordinate"] for c in record["cells"]),
                            "; ".join(c["coordinate"] + "=" + c["value"] for c in record["cells"]),
                            "; ".join(c["coordinate"] + ":" + str(c["formula"] or c["comment"] or "") for c in record["cells"] if c["formula"] or c["comment"])]
                # Each range reports real source row numbers, not batch ordinal numbers.
                before = len(self.units)
                self.table_batches(records, headings, {"sheet": sheet.title, "sheet_state": sheet.sheet_state}, display,
                                   "xlsx_cells", extra={"header_candidates": header_rows, "merged_ranges": merged,
                                   "hidden_columns": [k for k, v in sheet.column_dimensions.items() if v.hidden]})
                for unit in self.units[before:]:
                    batch = unit["structured"]["records"]
                    if batch:
                        unit["locator"].update(row_start=batch[0]["row"], row_end=batch[-1]["row"])
                        coordinates = [c["coordinate"] for r in batch for c in r["cells"]]
                        unit["locator"]["cell_range"] = f"A{batch[0]['row']}:{openpyxl.utils.get_column_letter(sheet.max_column)}{batch[-1]['row']}"
                        header_text = "\n".join("; ".join(c["coordinate"] + "=" + c["raw"] for c in row if c["raw"]) for row in header_rows)
                        unit["content_md"] = "# Sheet: " + sheet.title + "\n\n## Original header candidates (not inferred)\n" + header_text + "\n\n" + unit["content_md"]
                        unit["unit_id"] = ident("unit", [unit["document_id"], unit["locator"], unit["content_md"], unit["structured"]])
                    unit["parse_warnings"].extend(md_warning)
                    if any(c["formula"] and not c["cached_value"] for r in batch for c in r["cells"]):
                        unit["parse_warnings"].append("formula_cache_missing: do not invent calculated values")
        finally:
            formulas.close()
            cached.close()

    def pptx(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        presentation = Presentation(self.path)
        for si, slide in enumerate(presentation.slides, 1):
            shapes, assets, warnings = [], [], []
            def visit(shape, group=None):
                item = {"shape_id": shape.shape_id, "name": shape.name, "group_shape_id": group,
                        "bbox_emu": [shape.left, shape.top, shape.width, shape.height]}
                if shape.has_text_frame:
                    item["text"] = clean(shape.text)
                if shape.has_table:
                    item["table"] = [[c.text for c in row.cells] for row in shape.table.rows]
                if shape.has_chart:
                    try:
                        chart = shape.chart
                        item["chart"] = {"series": [{"name": s.name, "values": [cell_text(v) for v in s.values]} for s in chart.series],
                                         "plots": [{"categories": [str(c.label) for c in plot.categories]} for plot in chart.plots]}
                    except Exception as exc:
                        warnings.append("chart_metadata_failed: " + str(exc))
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    for child in shape.shapes:
                        visit(child, shape.shape_id)
                try:
                    if hasattr(shape, "image"):
                        image = shape.image
                        ref = self.asset(f"slide{si}_shape{shape.shape_id}.{image.ext}", image.blob)
                        item["asset_ref"] = ref
                        assets.append(ref)
                except (AttributeError, ValueError):
                    warnings.append(f"image_unavailable: shape {shape.shape_id}")
                # Preserve native connector endpoints; no invented causal interpretation.
                connectors = shape._element.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}stCxn") + shape._element.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}endCxn")
                if connectors:
                    item["connector_refs"] = [dict(c.attrib) for c in connectors]
                shapes.append(item)
            for shape in slide.shapes:
                visit(shape)
            notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None else ""
            blocks = []
            for item in shapes:
                if item.get("text", "").strip():
                    blocks.append(item["text"].strip())
                if item.get("table"):
                    blocks.append(md_table(item["table"]))
                if item.get("chart"):
                    blocks.append("Chart: " + canonical(item["chart"]))
            if notes.strip():
                blocks.append("Notes: " + clean(notes).strip())
            content = "\n\n".join(blocks)
            if assets:
                warnings.append("embedded_images_require_host_visual_review")
            self.add(content or "[Blank slide]", {"slide": si}, assets=assets, warnings=warnings,
                     structured={"kind": "pptx_shapes", "shapes": shapes, "notes": notes})

    def pdf(self):
        import pdfplumber
        with pdfplumber.open(self.path) as pdf:
            previous_page_schemas = []
            for pi, page in enumerate(pdf.pages, 1):
                tables = page.find_tables()
                raw_words = page.extract_words(x_tolerance=1, y_tolerance=2,
                                               keep_blank_chars=False, use_text_flow=False)
                text = clean(page.extract_text(layout=True) or "").strip()
                warnings, assets = [], []
                # Render only when the text layer is weak or the page contains raster imagery.
                # Text-heavy pages remain text-only to keep parsing cheap.
                has_images = bool(getattr(page, "images", None))
                needs_visual = (len(text) < 80) or has_images
                if needs_visual:
                    try:
                        image = page.to_image(resolution=110).original
                        buffer = io.BytesIO()
                        image.save(buffer, format="PNG")
                        assets.append(self.asset(f"page_{pi}.png", buffer.getvalue()))
                    except Exception as exc:
                        warnings.append("page_render_failed: " + str(exc))
                if not text:
                    warnings.append("no_text_layer: host_visual_extraction_required")
                elif not needs_visual:
                    warnings.append("page_image_skipped_text_layer_sufficient")
                words = [{"text": w["text"], "bbox": [str(w[k]) for k in ("x0", "top", "x1", "bottom")]} for w in raw_words]
                self.add(text or "[No PDF text layer; inspect supplied page image]", {"page": pi, "bbox": ["0", "0", str(page.width), str(page.height)]},
                         assets=assets, warnings=warnings, structured={"kind": "pdf_page", "words": words})
                current_page_schemas = []
                for ti, table in enumerate(tables, 1):
                    # Equal widths are common across unrelated tables in one document.
                    # A continuation must connect the previous page's bottom table to
                    # the next page's first table near the top of its body.
                    previous = next((schema for schema in sorted(previous_page_schemas,
                                     key=lambda s: s['bbox'][3], reverse=True)
                                     if ti == 1 and float(table.bbox[1]) <= float(page.height) * 0.15
                                     and float(schema['bbox'][3]) >= float(schema['page_height']) * 0.80
                                     and _pdf_same_x_span(schema.get("bbox"), table.bbox)), None)
                    header, data_rows, schema = _pdf_reconstruct_table(raw_words, table, previous, pi)
                    if not header and not data_rows:
                        continue
                    inherited_from = schema.get("inherited_from_page") if schema else None
                    if schema and len(header) >= 2:
                        schema['page_height'] = float(page.height)
                        current_page_schemas.append(schema)
                    self.table_batches(data_rows, header, {"page": pi, "table": ti, "bbox": [str(v) for v in table.bbox]},
                        lambda row: row, "pdf_table", first_record=1 if inherited_from else 2,
                        extra={"header": header, "schema_inherited_from_page": inherited_from,
                               "blank_policy": "preserve cells; inherit adjacent-page x-aligned schema only"})
                    for unit in self.units:
                        if unit["locator"].get("page") == pi and unit["locator"].get("table") == ti:
                            unit["asset_refs"] = assets
                previous_page_schemas = current_page_schemas
                if not tables:
                    self.units[-1]["parse_warnings"].append("no_ruled_table_detected: inspect text for borderless tables")

    def image(self):
        from PIL import Image
        with Image.open(self.path) as image:
            width, height = image.size
        ref = self.asset("original" + self.path.suffix.lower(), self.path.read_bytes())
        self.add("[Image: host must inspect original, labels, connectors and legend]",
                 {"image": self.path.name, "width": width, "height": height}, assets=[ref],
                 warnings=["host_visual_extraction_required"])

    def run(self):
        suffix = self.path.suffix.lower()
        if suffix == ".docx":
            self.docx()
        elif suffix == ".xlsx":
            self.xlsx()
        elif suffix == ".pptx":
            self.pptx()
        elif suffix == ".pdf":
            self.pdf()
        elif suffix in {".md", ".txt"}:
            self.add(self.path.read_text(encoding="utf-8-sig"), {"section": "document"})
        else:
            self.image()
        if not self.units:
            self.add("[Empty document]", {"section": "document"}, warnings=["empty_document"])
        return self.units


def parse_manifest(manifest_path, work, config=None, fresh=False):
    manifest_path, work = Path(manifest_path).resolve(), Path(work).resolve()
    manifest = read_json(manifest_path)
    require(isinstance(manifest, dict) and isinstance(manifest.get("files"), list) and bool(manifest["files"]), "Manifest needs nonempty files array")
    root = (manifest_path.parent / manifest.get("root", ".")).resolve()
    config = {**CONFIG, **(config or {})}
    require(config["max_chars"] >= 500 and config["max_rows"] >= 1, "Invalid parsing budgets")
    documents, paths, seen = [], {}, set()
    for item in manifest["files"]:
        require(isinstance(item, dict) and isinstance(item.get("path"), str), "Manifest entry needs path")
        path = (root / item["path"]).resolve()
        require(path.is_file(), "Missing input: " + str(path))
        require(path not in seen, "Duplicate input file")
        require(not path.is_relative_to(work), "Work directory must not contain input files")
        seen.add(path)
        role = item.get("role", "corpus")
        require(role in {"corpus", "questions", "reference"}, "Unknown input role")
        require(role != "corpus" or path.suffix.lower() in SUPPORTED, "Unsupported corpus format")
        try:
            relative = path.relative_to(root).as_posix()
        except ValueError as exc:
            raise ContractError("Input path escapes manifest root") from exc
        sha = hash_file(path)
        did = ident("doc", [relative, sha])
        documents.append({"document_id": did, "path": relative, "sha256": sha,
                          "format": path.suffix.lower().lstrip("."), "role": role, "parse_status": "excluded" if role != "corpus" else "pending"})
        paths[did] = path
    versions = dependency_versions()
    signature = engine_signature()
    fingerprint = ident("run", [ENGINE_VERSION, signature, documents, config, versions])
    state_path = work / "state.json"
    if state_path.exists():
        previous = read_json(state_path)
        if previous.get("fingerprint") == fingerprint and not fresh:
            return previous, True
        # Archive checkpoint JSON only; original input files are never moved/deleted.
        archive = work / "archive" / ident("checkpoint", [previous.get("fingerprint"), previous.get("audit"), previous.get("extractions")])
        archive.mkdir(parents=True, exist_ok=True)
        for name in ("state.json", "00_document_units.json", "00_extraction_plan.json", "01_mentions.json", "02_entity_map.json", "03_canonical_facts.json", "04_knowledge_graph.json"):
            existing = work / name
            if existing.is_file():
                shutil.copy2(existing, archive / name)
                existing.unlink()
    units = []
    work.mkdir(parents=True, exist_ok=True)
    for document in documents:
        if document["role"] != "corpus":
            continue
        adapter = Adapter(document, paths[document["document_id"]], work, config)
        try:
            units.extend(adapter.run())
            document["parse_status"] = "parsed"
        except Exception as exc:
            document["parse_status"] = "failed"
            document["parse_error"] = f"{type(exc).__name__}: {exc}"
            # Partial content isn't silently promoted to successful parsing.
            units.extend(adapter.units)
    state = {"schema_version": SCHEMA_VERSION, "engine_version": ENGINE_VERSION, "engine_signature": signature,
             "fingerprint": fingerprint, "config": config, "parser_versions": versions,
             "source_root": str(root), "documents": documents, "units": units,
             "extractions": {}, "alignment": empty_alignment(), "audit": [{"action": "parse", "fingerprint": fingerprint}]}
    write_json(work / "00_document_units.json", {"fingerprint": fingerprint, "documents": documents, "units": units})
    save_state(work, state)
    return state, False
