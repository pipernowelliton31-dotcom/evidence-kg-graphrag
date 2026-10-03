import importlib
import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import parse_documents as native


class MarkItDownParserTests(unittest.TestCase):
    def setUp(self):
        self.assertTrue((SCRIPTS / 'parse_markitdown.py').is_file(), 'Optional parser entry is missing')
        if importlib.util.find_spec('markitdown') is None:
            self.skipTest('Install optional requirements-markitdown.txt to test this route')
        self.optional = importlib.import_module('parse_markitdown')
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.inputs = self.root / 'inputs'
        self.inputs.mkdir()
        self.work = self.root / 'work'

    def manifest(self, names):
        path = self.root / 'manifest.json'
        path.write_text(json.dumps({'root': 'inputs', 'files': [
            {'path': name, 'role': 'questions' if name == 'questions.txt' else 'corpus'}
            for name in names]}), encoding='utf-8')
        return path

    def test_xlsx_ledger_matches_native_and_markdown_is_supplementary(self):
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        wb = Workbook()
        sheet = wb.active
        sheet.title = 'Ledger'
        sheet.append(['ID', 'Amount', 'Note'])
        sheet.append(['ITEM-7', '=2+3', 'hidden detail'])
        sheet['C2'].comment = Comment('source annotation', 'author')
        sheet.column_dimensions['C'].hidden = True
        sheet.row_dimensions[2].hidden = True
        sheet.merge_cells('A4:C4')
        sheet['A4'] = 'merged footer'
        wb.save(self.inputs / 'ledger.xlsx')
        manifest = self.manifest(['ledger.xlsx'])
        original, _ = native.parse_manifest(manifest, self.root / 'native')
        result, _ = self.optional.parse_manifest(manifest, self.work)
        self.assertEqual(original['units'], result['units'])
        self.assertEqual(result['documents'][0]['parse_status'], 'parsed')
        md = self.work / result['documents'][0]['markdown_ref']
        self.assertIn('ITEM-7', md.read_text(encoding='utf-8'))
        self.assertIn('markitdown', result['parser_versions'])

    def test_pptx_keeps_slide_structure_and_notes_with_markdown_reading_view(self):
        from pptx import Presentation
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = 'Release event'
        slide.placeholders[1].text = 'A new feature shipped'
        slide.notes_slide.notes_text_frame.text = 'Source note must survive'
        prs.save(self.inputs / 'slides.pptx')
        manifest = self.manifest(['slides.pptx'])
        original, _ = native.parse_manifest(manifest, self.root / 'native')
        result, _ = self.optional.parse_manifest(manifest, self.work)
        self.assertEqual(result['units'][0]['structured'], original['units'][0]['structured'])
        self.assertEqual(result['units'][0]['locator']['slide'], 1)
        self.assertIn('Release event', result['units'][0]['content_md'])
        self.assertIn('Source note must survive', result['units'][0]['content_md'])
        self.assertTrue((self.work / result['documents'][0]['markdown_ref']).is_file())

    def test_cache_reuse_source_invalidation_and_native_backend_separation(self):
        (self.inputs / 'notes.txt').write_text('A source fact', encoding='utf-8')
        (self.inputs / 'questions.txt').write_text('Do not extract this answer', encoding='utf-8')
        manifest = self.manifest(['notes.txt', 'questions.txt'])
        first, cached = self.optional.parse_manifest(manifest, self.work)
        self.assertFalse(cached)
        second, cached = self.optional.parse_manifest(manifest, self.work)
        self.assertTrue(cached)
        self.assertEqual(first['fingerprint'], second['fingerprint'])
        self.assertEqual(len(first['units']), 1)
        self.assertEqual(first['documents'][1]['parse_status'], 'excluded')
        (self.inputs / 'notes.txt').write_text('A changed source fact', encoding='utf-8')
        changed, cached = self.optional.parse_manifest(manifest, self.work)
        self.assertFalse(cached)
        self.assertNotEqual(first['fingerprint'], changed['fingerprint'])
        native_state, cached = native.parse_manifest(manifest, self.work)
        self.assertFalse(cached)
        self.assertNotEqual(native_state['fingerprint'], changed['fingerprint'])
        self.assertTrue(list((self.work / 'archive').iterdir()))

    def test_conversion_failure_is_visible_and_not_reported_as_success(self):
        from pptx import Presentation
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = 'Native text survives'
        prs.save(self.inputs / 'broken.pptx')
        with patch.object(self.optional.MarkdownAdapter, 'markdown', side_effect=RuntimeError('conversion failed')):
            result, _ = self.optional.parse_manifest(self.manifest(['broken.pptx']), self.work)
        self.assertEqual(result['documents'][0]['parse_status'], 'failed')
        self.assertIn('conversion failed', result['documents'][0]['parse_error'])

    def test_cli_same_contract_and_missing_dependency_is_clear(self):
        (self.inputs / 'notes.txt').write_text('A source fact', encoding='utf-8')
        manifest = self.manifest(['notes.txt'])
        proc = subprocess.run([sys.executable, str(SCRIPTS / 'parse_markitdown.py'),
                               '--manifest', str(manifest), '--work', str(self.work), '--fresh'],
                              capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(proc.returncode, 0, proc.stderr + proc.stdout)
        result = json.loads(proc.stdout)
        self.assertTrue(result['ok'])
        self.assertEqual(result['parser_backend'], 'markitdown')
        with patch.object(self.optional, 'require_converter', side_effect=self.optional.ContractError('Install requirements-markitdown.txt')):
            with self.assertRaisesRegex(self.optional.ContractError, 'requirements-markitdown.txt'):
                self.optional.parse_manifest(manifest, self.root / 'missing-dependency')


class MissingOptionalDependencyTests(unittest.TestCase):
    def test_cli_reports_missing_dependency_without_creating_work(self):
        with tempfile.TemporaryDirectory() as root:
            work = Path(root) / 'work'
            proc = subprocess.run([sys.executable, '-S', str(SCRIPTS / 'parse_markitdown.py'),
                                   '--manifest', str(Path(root) / 'inputs.json'), '--work', str(work)],
                                  capture_output=True, text=True, encoding='utf-8')
            self.assertEqual(proc.returncode, 2)
            self.assertIn('requirements-markitdown.txt', json.loads(proc.stdout)['error'])
            self.assertFalse(work.exists())


if __name__ == '__main__':
    unittest.main()
